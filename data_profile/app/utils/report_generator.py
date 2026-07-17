from io import BytesIO
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt


def _format_number(value):
    if isinstance(value, float):
        return f"{value:,.2f}"
    if isinstance(value, int):
        return f"{value:,}"
    return str(value)


def _format_percentage(value):
    return f"{value:.2f}%"


def _build_executive_summary(profile):
    row_count = profile.get("row_count", 0)
    column_count = profile.get("column_count", 0)
    duplicate_count = profile.get("duplicate_count", 0)
    quality_score = profile.get("quality_score", 0.0)
    nulls = profile.get("null_percentages", {})
    pk_candidates = profile.get("candidate_primary_keys", [])

    summary_parts = []
    summary_parts.append(
        f"This table contains {_format_number(row_count)} rows and {column_count} columns with an overall data quality score of {quality_score:.2f}/100."
    )

    if quality_score >= 90:
        summary_parts.append("The assessment indicates strong overall data quality.")
    elif quality_score >= 75:
        summary_parts.append(
            "The assessment indicates generally good data quality, with a few issues to address before broader operational use."
        )
    elif quality_score >= 50:
        summary_parts.append(
            "The assessment indicates moderate data quality concerns that should be investigated before downstream reporting."
        )
    else:
        summary_parts.append(
            "The assessment indicates significant data quality issues that require remediation before this table is used in production."
        )

    if duplicate_count == 0:
        summary_parts.append("No duplicate records were detected.")
    else:
        summary_parts.append(
            f"{_format_number(duplicate_count)} duplicate records were detected and should be investigated."
        )

    high_null_columns = [
        (column, percent)
        for column, percent in nulls.items()
        if 15.0 <= percent < 100.0
    ]
    high_null_columns.sort(key=lambda item: -item[1])
    if high_null_columns:
        formatted_cols = ", ".join(
            [f"{col} at {percent:.0f}%" for col, percent in high_null_columns[:5]]
        )
        summary_parts.append(
            f"Several columns have elevated null rates, including {formatted_cols}."
        )

    fully_null_columns = [column for column, percent in nulls.items() if percent == 100.0]
    if fully_null_columns:
        summary_parts.append(
            f"The following columns are fully null and should be reviewed or removed from reporting outputs: {', '.join(fully_null_columns)}."
        )

    if pk_candidates:
        summary_parts.append(
            f"{pk_candidates[0]} appears to be a strong primary key candidate."
        )

    return " ".join(summary_parts)


def create_word_report(profile):
    doc = Document()

    logo_path = Path(__file__).resolve().parents[1] / "Logo Sigmatic-r0d1-01.png"
    if logo_path.exists():
        try:
            doc.add_picture(str(logo_path), width=Inches(1.5))
        except Exception:
            pass

    title = doc.add_heading("Sigmatic Data Profiling Report", level=1)
    title.runs[0].font.size = Pt(24)

    if profile.get("generated_at") or profile.get("workspace_host") or profile.get("warehouse"):
        header = doc.add_paragraph()
        header.add_run("Report generated: ").bold = True
        header.add_run(profile.get("generated_at", "Unknown"))
        header.add_run("\nEnvironment: ").bold = True
        header.add_run(str(profile.get("workspace_host", "Unknown")))
        header.add_run("\nSQL Warehouse: ").bold = True
        header.add_run(str(profile.get("warehouse", "Unknown")))
        header.add_run("\nTable: ").bold = True
        header.add_run(
            f"{profile.get('catalog', 'unknown')}.{profile.get('schema_name', 'unknown')}.{profile.get('table_name', 'unknown')}"
        )

    doc.add_heading("Executive Summary", level=2)
    doc.add_paragraph(_build_executive_summary(profile))

    components = profile.get("quality_score_components", {})
    if components:
        doc.add_heading("Quality Score Breakdown", level=2)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Dimension"
        hdr[1].text = "Score"
        hdr[2].text = "Explanation"

        for dimension, label, explanation in [
            ("completeness", "Completeness", "Coverage of required values and null risk."),
            ("uniqueness", "Uniqueness", "Duplicate risk and candidate key quality."),
            ("validity", "Validity", "Data type and structural consistency."),
            ("governance", "Governance", "Fully null columns and schema risk."),
        ]:
            value = components.get(dimension)
            if value is None:
                continue
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = f"{value:.2f}"
            cells[2].text = explanation

        total_row = table.add_row().cells
        total_row[0].text = "Overall Score"
        total_row[1].text = f"{components.get('overall_score', 0.0):.2f}"
        total_row[2].text = "Weighted combination of all quality dimensions."

    doc.add_heading("Key Findings", level=2)
    doc.add_paragraph(_build_executive_summary(profile))

    doc.add_heading("Recommendations", level=2)
    for recommendation in profile.get("recommendations", []):
        doc.add_paragraph(recommendation, style="List Bullet")

    doc.add_heading("Schema", level=2)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"

    hdr = table.rows[0].cells
    hdr[0].text = "Column"
    hdr[1].text = "Type"
    hdr[2].text = "Nullable"
    hdr[3].text = "Position"

    for row in profile.get("schema", []):
        cells = table.add_row().cells
        cells[0].text = str(row.get("column_name", ""))
        cells[1].text = str(row.get("data_type", ""))
        cells[2].text = str(row.get("is_nullable", ""))
        cells[3].text = str(row.get("ordinal_position", ""))

    
    doc.add_page_break()

    doc.add_heading("Null Percentages", level=2)
    null_table = doc.add_table(rows=1, cols=2)
    null_table.style = "Table Grid"
    hdr = null_table.rows[0].cells
    hdr[0].text = "Column"
    hdr[1].text = "Null %"
    for col, value in sorted(profile.get("null_percentages", {}).items(), key=lambda item: -item[1]):
        cells = null_table.add_row().cells
        cells[0].text = col
        cells[1].text = _format_percentage(value)

    doc.add_heading("Cardinality", level=2)
    cardinality_table = doc.add_table(rows=1, cols=2)
    cardinality_table.style = "Table Grid"
    hdr = cardinality_table.rows[0].cells
    hdr[0].text = "Column"
    hdr[1].text = "Unique Count"
    for col, value in sorted(profile.get("unique_counts", {}).items(), key=lambda item: item[0]):
        cells = cardinality_table.add_row().cells
        cells[0].text = col
        cells[1].text = _format_number(value)

    numeric_stats = profile.get("numeric_statistics", [])
    if numeric_stats:
        doc.add_heading("Numeric Statistics", level=2)
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Column"
        hdr[1].text = "Min"
        hdr[2].text = "Max"
        hdr[3].text = "Average"
        hdr[4].text = "Sum"

        for row in numeric_stats:
            cells = table.add_row().cells
            cells[0].text = str(row.get("column_name", ""))
            cells[1].text = _format_number(row.get("min_value", ""))
            cells[2].text = _format_number(row.get("max_value", ""))
            cells[3].text = _format_number(row.get("avg_value", ""))
            cells[4].text = _format_number(row.get("sum_value", ""))

    doc.add_page_break()

    doc.add_heading("Candidate Primary Keys", level=2)
    if profile.get("candidate_primary_keys"):
        doc.add_paragraph(", ".join(profile.get("candidate_primary_keys", [])))
    else:
        doc.add_paragraph("No candidate primary keys identified.")

    doc.add_heading("Data Type Issues", level=2)
    data_type_issues = profile.get("data_type_issues", [])
    if data_type_issues:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Column"
        hdr[1].text = "Issue"
        hdr[2].text = "Severity"
        for issue in data_type_issues:
            cells = table.add_row().cells
            cells[0].text = str(issue.get("column_name", ""))
            cells[1].text = str(issue.get("issue", ""))
            cells[2].text = str(issue.get("severity", ""))
    else:
        doc.add_paragraph("No data type issues detected.")

    doc.add_heading("Quality Warnings", level=2)
    quality_warnings = profile.get("quality_warnings", [])
    if quality_warnings:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Issue"
        hdr[1].text = "Details"
        hdr[2].text = "Severity"
        for warning in quality_warnings:
            cells = table.add_row().cells
            cells[0].text = str(warning.get("warning_type", ""))
            cells[1].text = str(warning.get("details", ""))
            cells[2].text = str(warning.get("severity", ""))
    else:
        doc.add_paragraph("No high-severity quality warnings detected.")

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output